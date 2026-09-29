"""التقارير المالية الأسبوعية والشهرية والسنوية، والتصدير إلى Excel."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..domain.periods import RIYADH
from ..models import AllocationEvent, Entity, PeriodClosure, utcnow
from . import metrics as m

CALC = {"percentage": "نسبة", "fixed_amount": "مبلغ ثابت"}


@dataclass
class Report:
    entity: Entity
    kind: str
    snap: m.Snapshot
    rows: list[dict]  # معاملة لكل صف مع توزيعها على البنود
    closures: list[dict]

    @property
    def title(self) -> str:
        return f"تقرير {m.REPORT_KINDS[self.kind]} — {self.snap.range.label}"


def build_report(session: Session, entity: Entity, kind: str, anchor: date) -> Report:
    rng = m.report_range(kind, anchor)
    snap = m.snapshot(session, entity.id, rng)
    names = {b.bucket_id: b.name for b in snap.buckets}

    by_txn: dict[int, dict[int, int]] = {}
    if snap.transactions:
        for e in session.scalars(
            select(AllocationEvent).where(AllocationEvent.transaction_id.in_([t.id for t in snap.transactions]))
        ):
            by_txn.setdefault(e.transaction_id, {})[e.bucket_id] = e.allocated
    rows = [
        {"txn": t, "allocations": by_txn.get(t.id, {})}
        for t in snap.transactions
    ]

    closures = []
    for c in session.scalars(
        select(PeriodClosure)
        .where(PeriodClosure.entity_id == entity.id, PeriodClosure.voided.is_(False))
        .order_by(PeriodClosure.period_key)
    ):
        if c.period_key == "0000-00-00":
            continue
        start = date.fromisoformat(c.period_key)
        if rng.start <= start < rng.end:
            closures.append({"closure": c, "name": names.get(c.bucket_id, f"بند #{c.bucket_id}")})
    return Report(entity=entity, kind=kind, snap=snap, rows=rows, closures=closures)


# ------------------------------------------------------------------ Excel

_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="0F6E5A")
_TITLE = Font(bold=True, size=14)
_BOLD = Font(bold=True)
_MONEY = "#,##0.00"
_THIN = Border(bottom=Side(style="thin", color="E3E6EA"))


def _sheet(wb: Workbook, title: str, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def _header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font, cell.fill = _HEAD, _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _money(cell) -> None:
    cell.number_format = _MONEY


def to_excel(report: Report) -> bytes:
    wb = Workbook()
    snap, entity = report.snap, report.entity

    # --- الملخص
    ws = _sheet(wb, "الملخص", first=True)
    ws["A1"] = report.title
    ws["A1"].font = _TITLE
    info = [
        ("الشركة", entity.name),
        ("السجل التجاري", entity.cr_number),
        ("الرقم الضريبي", entity.vat_number),
        ("الفترة", f"{snap.range.start.isoformat()} إلى {(snap.range.end - timedelta(days=1)).isoformat()}"),
        ("تاريخ الإصدار", utcnow().astimezone(RIYADH).strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (label, value) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=label).font = _BOLD
        ws.cell(row=i, column=2, value=value)

    totals = [
        ("إجمالي الوارد", snap.incoming),
        ("المخصص للبنود", snap.allocated),
        ("الفائض", snap.surplus),
        ("عجز الفترات المقفلة", snap.closed_deficit),
        ("العجز الحالي (وقت الإصدار)", snap.deficit),
    ]
    start = len(info) + 4
    ws.cell(row=start, column=1, value="الإجماليات (ر.س)").font = _BOLD
    for i, (label, value) in enumerate(totals, start + 1):
        ws.cell(row=i, column=1, value=label)
        _money(ws.cell(row=i, column=2, value=value / 100))
    ws.cell(row=start + len(totals) + 1, column=1, value="عدد المعاملات")
    ws.cell(row=start + len(totals) + 1, column=2, value=snap.count)

    table = start + len(totals) + 3
    _header(ws, table, ["الأولوية", "البند", "النوع", "القيمة", "المخصص خلال الفترة", "الهدف", "المتجمّع حاليًا", "العجز الحالي"])
    for r, b in enumerate(snap.buckets, table + 1):
        value = f"{b.value / 100:g}%" if b.calc_type and b.calc_type.value == "percentage" else b.value / 100
        cells = [b.priority, b.name, CALC.get(b.calc_type.value if b.calc_type else "", ""), value, b.allocated / 100, b.target / 100, b.funded / 100, b.deficit / 100]
        for c, v in enumerate(cells, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = _THIN
            if c >= 5 or (c == 4 and isinstance(v, float)):
                _money(cell)
    _widths(ws, [26, 26, 12, 14, 20, 14, 18, 16])

    # --- المعاملات
    ws = _sheet(wb, "المعاملات")
    buckets = snap.buckets
    _header(ws, 1, ["#", "تاريخ الإيداع", "النوع", "المبلغ", *[b.name for b in buckets], "الفائض", "ملاحظة"])
    for r, row in enumerate(report.rows, 2):
        t = row["txn"]
        cells = [
            t.id,
            t.occurred_at.astimezone(RIYADH).strftime("%Y-%m-%d %H:%M"),
            f"قيد عكسي لـ#{t.reversal_of}" if t.reversal_of else "وارد",
            t.amount / 100,
            *[row["allocations"].get(b.bucket_id, 0) / 100 for b in buckets],
            t.surplus / 100,
            t.note,
        ]
        for c, v in enumerate(cells, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = _THIN
            if 4 <= c <= 5 + len(buckets):
                _money(cell)
    total_row = len(report.rows) + 2
    ws.cell(row=total_row, column=3, value="الإجمالي").font = _BOLD
    for c in range(4, 6 + len(buckets)):
        col = get_column_letter(c)
        cell = ws.cell(row=total_row, column=c, value=f"=SUM({col}2:{col}{total_row - 1})" if report.rows else 0)
        cell.font = _BOLD
        _money(cell)
    _widths(ws, [8, 18, 16, 14, *[16] * len(buckets), 14, 30])
    ws.freeze_panes = "B2"

    # --- إقفال الفترات
    ws = _sheet(wb, "إقفال الفترات")
    _header(ws, 1, ["البند", "بداية الفترة", "الهدف", "المموّل", "العجز النهائي"])
    for r, item in enumerate(report.closures, 2):
        c = item["closure"]
        for col, v in enumerate([item["name"], c.period_key, c.target / 100, c.funded / 100, c.closed_deficit / 100], 1):
            cell = ws.cell(row=r, column=col, value=v)
            if col >= 3:
                _money(cell)
    _widths(ws, [24, 14, 14, 14, 16])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
