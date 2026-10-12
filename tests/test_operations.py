"""Transactions, reversals, metrics, dashboard and reports on a company with an active policy."""
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import models
from app.domain import formula
from app.domain.periods import RIYADH
from app.models import CalcType, Frequency, Role
from app.services import companies, dashboards, metrics
from app.services import policies as pol
from app.services import transactions as txs

SAR = 100


def at(y, mo, d, h=12):
    return datetime(y, mo, d, h, tzinfo=RIYADH)


@pytest.fixture
def setup(db, make_user):
    accountant = make_user("acc")
    entity = companies.register_company(
        db, accountant, companies.CompanyInput(name="مؤسسة النخيل"), companies.ReviewerInput("سارة", "sara", "Password-123!")
    )
    reviewer = db.query(models.User).filter_by(username="sara").one()
    reviewer.must_change_password = False
    draft = pol.start_draft(db, entity, accountant)
    pol.save_draft(
        db,
        draft,
        [
            pol.BucketInput(None, "الضريبة", 1, CalcType.PERCENTAGE, 1500, Frequency.IMMEDIATE, None, "", True, new_destination_iban="SA1100000000000000000001"),
            pol.BucketInput(None, "الإيجار", 2, CalcType.FIXED_AMOUNT, 10_000 * SAR, Frequency.IMMEDIATE, None, "حساب الإيجار", False, new_destination_iban="SA1100000000000000000003"),
            pol.BucketInput(None, "الأرباح", 3, CalcType.PERCENTAGE, 500, Frequency.MONTHLY, None, "", False, new_destination_iban="SA1100000000000000000002"),
        ],
        "",
        accountant,
    )
    pol.submit(db, draft, accountant, "")
    pol.approve(db, draft, reviewer)
    db.commit()
    ids = {pb.name: pb.bucket_id for pb in draft.buckets}
    return entity, accountant, ids


def record(db, entity, user, amount, when, **kw):
    txn = txs.record_transaction(db, entity, user, amount * SAR, when, **kw)
    db.commit()
    return txn


def states(db, entity):
    return {s.bucket_id: (s.period_key, s.funded, s.carried_deficit) for s in db.query(models.BucketPeriodState).filter_by(entity_id=entity.id)}


def test_record_matches_engine_scenario(db, setup):
    entity, user, ids = setup
    t1 = record(db, entity, user, 4_000, at(2026, 9, 2))
    t2 = record(db, entity, user, 9_000, at(2026, 9, 10))
    alloc = lambda t: {e.bucket_id: e.allocated for e in t.events}  # noqa: E731
    assert alloc(t1) == {ids["الضريبة"]: 600 * SAR, ids["الإيجار"]: 3_400 * SAR, ids["الأرباح"]: 0}
    assert alloc(t2)[ids["الأرباح"]] == 650 * SAR  # 450 + 200 carried
    assert (t1.surplus, t2.surplus) == (0, 400 * SAR)
    assert states(db, entity)[ids["الإيجار"]] == ("0000-00-00", 10_000 * SAR, 0)  # no due date: perpetual period


def test_occurred_at_round_trips_with_timezone(db, setup):
    entity, user, _ = setup
    txn = record(db, entity, user, 100, at(2026, 9, 29, 1))  # 1 AM Riyadh = 22:00 UTC the day before
    db.expire_all()
    stored = db.get(models.Transaction, txn.id).occurred_at
    assert stored.astimezone(RIYADH) == at(2026, 9, 29, 1)


def test_requires_active_policy(db, make_user):
    user = make_user("acc2")
    entity = companies.register_company(db, user, companies.CompanyInput(name="بدون سياسة"), companies.ReviewerInput("ر", "rev2", "Password-123!"))
    with pytest.raises(txs.TransactionError, match="سياسة نشطة"):
        txs.record_transaction(db, entity, user, 100, at(2026, 9, 1))


def test_closed_period_rejected(db, setup):
    entity, user, _ = setup
    record(db, entity, user, 1_000, at(2026, 10, 2))
    with pytest.raises(txs.TransactionError, match="مقفلة"):
        txs.record_transaction(db, entity, user, 100, at(2026, 9, 30))


def test_external_ref_is_idempotent(db, setup):
    entity, user, _ = setup
    a = record(db, entity, user, 1_000, at(2026, 9, 1), external_ref="bank-123")
    b = record(db, entity, user, 1_000, at(2026, 9, 1), external_ref="bank-123")
    assert a.id == b.id
    assert db.query(models.Transaction).count() == 1


@pytest.fixture
def no_transfers(monkeypatch):
    """A reversal is only possible before money leaves for a beneficiary, so this tests reversal logic without transfers."""
    monkeypatch.setattr(txs.payouts, "pay_on_allocation", lambda *a, **k: [])
    monkeypatch.setattr(txs.payouts, "run_due", lambda *a, **k: [])


def test_reversal_restores_state_and_only_latest(db, setup, no_transfers):
    entity, user, ids = setup
    record(db, entity, user, 4_000, at(2026, 9, 2))
    before = states(db, entity)
    t2 = record(db, entity, user, 9_000, at(2026, 9, 10))
    t1 = db.query(models.Transaction).order_by(models.Transaction.id).first()

    with pytest.raises(txs.TransactionError, match="آخر معاملة"):
        txs.reverse_transaction(db, t1, user, "غلط")
    with pytest.raises(txs.TransactionError, match="سبب"):
        txs.reverse_transaction(db, t2, user, "")

    rev = txs.reverse_transaction(db, t2, user, "مبلغ مكرر")
    db.commit()
    assert states(db, entity) == before
    assert (rev.amount, rev.surplus) == (-t2.amount, -t2.surplus)
    # After reversing the second, the first is now the latest and can be reversed; the state returns to empty
    assert txs.reversible_transaction(db, entity.id).id == t1.id
    txs.reverse_transaction(db, t1, user, "تجربة")
    db.commit()
    assert states(db, entity) == {}


def test_reversal_voids_period_closures(db, setup, no_transfers):
    entity, user, _ = setup
    record(db, entity, user, 1_000, at(2026, 9, 5))
    t2 = record(db, entity, user, 1_000, at(2026, 10, 5))
    assert db.query(models.PeriodClosure).filter_by(voided=False).count() == 1  # the profit month
    txs.reverse_transaction(db, t2, user, "تاريخ غلط")
    db.commit()
    assert db.query(models.PeriodClosure).filter_by(voided=False).count() == 0
    assert {s[0] for s in states(db, entity).values()} == {"0000-00-00", "2026-09-01"}


def test_snapshot_and_variables(db, setup):
    entity, user, ids = setup
    record(db, entity, user, 500, at(2026, 8, 20))  # outside the month
    record(db, entity, user, 4_000, at(2026, 9, 2))
    record(db, entity, user, 9_000, at(2026, 9, 10))

    snap = metrics.snapshot(db, entity.id, metrics.resolve_range("this_month", date(2026, 9, 15)), today=date(2026, 9, 15))
    # The August deposit funded 425 of rent (no due date = perpetual period), so September only adds 9,575 and the surplus is larger
    assert (snap.incoming, snap.surplus, snap.count) == (13_000 * SAR, 825 * SAR, 2)
    rent = next(b for b in snap.buckets if b.name == "الإيجار")
    assert (rent.allocated, rent.funded, rent.deficit) == (9_575 * SAR, 10_000 * SAR, 0)
    values = metrics.variables(snap)
    assert values["الفائض"] == 825 and values["مخصص_الضريبة"] == 1_950
    assert formula.evaluate("الفائض / الوارد * 100", values) == pytest.approx(825 / 13_000 * 100)


def test_formula_safety():
    ok = {"الفائض": 10.0, "الوارد": 0.0}
    assert formula.evaluate("max(الفائض, 5) × ٢", ok) == 20
    assert formula.evaluate("الفائض / الوارد", ok) is None  # division by zero
    for bad in ["__import__('os')", "الفائض.real", "[1,2]", "الفائض ** 2", "غير_موجود + 1", "lambda: 1"]:
        with pytest.raises(formula.FormulaError):
            formula.evaluate(bad, ok)


# ------------------------------------------------------------------ through the UI


def login(client, username, password="password123"):
    client.post("/login", data={"username": username, "password": password})


def test_transaction_pages_and_roles(client, db, setup):
    entity, _, _ = setup
    login(client, "acc")
    response = client.post(
        f"/entities/{entity.id}/transactions",
        data={"amount": "4,000", "occurred_at": "2026-09-02T10:00", "note": "دفعة"},
        follow_redirects=False,
    )
    assert response.status_code == 303 and response.headers["location"].startswith("/transactions/")
    detail = client.get(response.headers["location"]).text
    assert "600.00" in detail and "مموَّل جزئيًا" in detail

    client.post("/logout")
    login(client, "sara", "Password-123!")
    # Reviewers only review the policy: every other company page is closed to them
    for page in ("transactions", "dashboard", "accounts", "balances", "reports", "audit", "settings"):
        assert client.get(f"/entities/{entity.id}/{page}").status_code == 403, page
    assert client.get(f"/entities/{entity.id}/policy").status_code == 200
    assert client.post(f"/entities/{entity.id}/transactions", data={"amount": "1", "occurred_at": "2026-09-02T10:00"}).status_code == 403


def test_future_date_rejected(client, db, setup):
    entity, _, _ = setup
    login(client, "acc")
    client.post(f"/entities/{entity.id}/transactions", data={"amount": "100", "occurred_at": "2099-01-01T10:00"})
    assert db.query(models.Transaction).count() == 0


def test_dashboard_default_widgets_and_customization(client, db, setup):
    entity, user, _ = setup
    record(db, entity, user, 4_000, datetime.now(RIYADH))
    login(client, "acc")
    page = client.get(f"/entities/{entity.id}/dashboard").text
    assert "الوارد هذا الشهر" in page and "4,000.00" in page and "تقدّم البنود" in page

    dash = db.query(models.Dashboard).filter_by(entity_id=entity.id).one()
    assert len(dash.widgets) == len(dashboards.DEFAULT_LAYOUT)

    # A custom metric, then a card showing it
    client.post(f"/entities/{entity.id}/metrics", data={"name": "نسبة التخصيص", "formula": "(مخصص_الضريبة + مخصص_الإيجار) / الوارد * 100", "fmt": "percent"})
    metric = db.query(models.CustomMetric).one()
    client.post(f"/entities/{entity.id}/dashboard/widgets", data={"kind": "kpi", "title": "نسبة التخصيص", "metric": f"custom:{metric.id}", "period": "this_month"})
    db.expire_all()
    assert len(dash.widgets) == len(dashboards.DEFAULT_LAYOUT) + 1
    assert "100.0%" in client.get(f"/entities/{entity.id}/dashboard").text  # (600 + 3,400) / 4,000

    # An unsafe formula is not saved
    client.post(f"/entities/{entity.id}/metrics", data={"name": "خطر", "formula": "__import__('os')", "fmt": "number"})
    assert db.query(models.CustomMetric).count() == 1

    # Saving the layout
    widget = dash.widgets[0]
    assert client.post(f"/entities/{entity.id}/dashboard/layout", json={"items": [{"id": widget.id, "x": 6, "y": 5, "w": 4, "h": 3}]}).json() == {"ok": True}
    db.refresh(widget)
    assert (widget.x, widget.y, widget.w, widget.h) == (6, 5, 4, 3)

    # Deleting the metric deletes the card showing it
    client.post(f"/metrics/{metric.id}/delete")
    db.expire_all()
    assert len(dash.widgets) == len(dashboards.DEFAULT_LAYOUT)


def test_reviewer_cannot_edit_dashboard(client, db, setup):
    entity, _, _ = setup
    login(client, "sara", "Password-123!")
    page = client.get(f"/entities/{entity.id}/dashboard").text
    assert "إضافة عنصر" not in page
    assert client.post(f"/entities/{entity.id}/dashboard/widgets", data={"kind": "kpi", "metric": "surplus"}).status_code == 403


def test_reports_page_and_excel(client, db, setup):
    entity, user, _ = setup
    record(db, entity, user, 4_000, at(2026, 9, 2))
    record(db, entity, user, 9_000, at(2026, 9, 10))
    login(client, "acc")
    page = client.get(f"/entities/{entity.id}/reports?kind=month&anchor=2026-09-15").text
    assert "13,000.00" in page and "400.00" in page

    response = client.get(f"/entities/{entity.id}/reports.xlsx?kind=month&anchor=2026-09-15")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["الملخص", "المعاملات", "إقفال الفترات"]
    txn_sheet = wb["المعاملات"]
    assert txn_sheet.sheet_view.rightToLeft
    assert [txn_sheet.cell(row=r, column=4).value for r in (2, 3)] == [4000, 9000]

    week = client.get(f"/entities/{entity.id}/reports?kind=week&anchor=2026-09-02").text
    assert "4,000.00" in week and "9,000.00" not in week

    assert "تنزيل تقرير" in client.get(f"/entities/{entity.id}/audit").text


def test_report_range_boundaries():
    assert metrics.report_range("week", date(2026, 9, 29)).start == date(2026, 9, 27)
    assert metrics.report_range("month", date(2026, 12, 15)).end == date(2027, 1, 1)
    assert metrics.report_range("year", date(2026, 5, 1)).start == date(2026, 1, 1)
